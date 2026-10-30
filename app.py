
from flask import Flask, request, redirect, session, render_template_string, flash
from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash, check_password_hash
from pathlib import Path
from datetime import datetime

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY")
CODIGO_ADMIN = os.environ.get("CODIGO_ADMIN")
ARCHIVO = Path(__file__).parent / "registros_itnh.xlsx"


def preparar_excel():
    if not ARCHIVO.exists():
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Estudiantes"
        hoja.append(["ID", "Usuario", "Correo", "Contraseña", "Fecha"])

        notas = libro.create_sheet("Notas")
        notas.append(["Correo", "Materia", "Período 1", "Período 2", "Período 3"])
        libro.save(ARCHIVO)
        libro.close()
    else:
        libro = load_workbook(ARCHIVO)

        if "Notas" not in libro.sheetnames:
            notas = libro.create_sheet("Notas")
            notas.append(["Correo", "Materia", "Período 1", "Período 2", "Período 3"])

        libro.save(ARCHIVO)
        libro.close()


def buscar_usuario(correo):
    preparar_excel()
    libro = load_workbook(ARCHIVO)
    hoja = libro["Estudiantes"]

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[2] and fila[2].lower() == correo.lower():
            libro.close()
            return fila

    libro.close()
    return None


def obtener_notas(correo):
    preparar_excel()
    libro = load_workbook(ARCHIVO)
    hoja = libro["Notas"]

    notas = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila[0] and str(fila[0]).lower() == correo.lower():
            notas.append(fila[1:])

    libro.close()
    return notas


HTML = """
<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Instituto Técnico Nuevo Horizonte</title>

<style>
* {
    box-sizing: border-box;
}

body {
    margin: 0;
    font-family: Arial, sans-serif;
    background: #f1f5f9;
    color: #1e293b;
}

header {
    background: #123b65;
    color: white;
    padding: 25px;
    text-align: center;
}

header h1 {
    margin: 0;
}

nav {
    background: #0b2948;
    padding: 15px;
    text-align: center;
}

nav a {
    color: white;
    text-decoration: none;
    margin: 10px;
    font-weight: bold;
}

nav a:hover {
    color: #facc15;
}

.contenedor {
    max-width: 850px;
    margin: 35px auto;
    background: white;
    padding: 30px;
    border-radius: 12px;
    box-shadow: 0 3px 12px #00000015;
}

h2 {
    color: #123b65;
}

input {
    display: block;
    width: 100%;
    padding: 12px;
    margin: 12px 0;
    border: 1px solid #cbd5e1;
    border-radius: 6px;
}

button {
    background: #123b65;
    color: white;
    border: none;
    padding: 12px 22px;
    border-radius: 6px;
    cursor: pointer;
}

button:hover {
    background: #1d568c;
}

table {
    width: 100%;
    border-collapse: collapse;
    margin-top: 20px;
}

th, td {
    border: 1px solid #cbd5e1;
    padding: 12px;
    text-align: center;
}

th {
    background: #123b65;
    color: white;
}

.mensaje {
    background: #e0f2fe;
    padding: 12px;
    margin-bottom: 15px;
    border-radius: 6px;
}

footer {
    text-align: center;
    padding: 20px;
    color: #64748b;
}
</style>
</head>

<body>

<header>
    <h1>Instituto Técnico Nuevo Horizonte</h1>
    <p>Educación, conocimiento y futuro</p>
</header>

<nav>
    <a href="/">Inicio</a>
    {% if session.get("usuario") %}
        <a href="/notas">Mis Notas</a>
        <a href="/salir">Cerrar sesión</a>
    {% else %}
        <a href="/registro">Registrarse</a>
        <a href="/login">Iniciar sesión</a>
    {% endif %}
</nav>

<div class="contenedor">

{% with mensajes = get_flashed_messages() %}
    {% for mensaje in mensajes %}
        <div class="mensaje">{{ mensaje }}</div>
    {% endfor %}
{% endwith %}

{% if pagina == "inicio" %}
    <h2>Bienvenidos</h2>
    <p>
        Bienvenidos al Instituto Técnico Nuevo Horizonte.
        Este espacio permite a los estudiantes registrarse,
        iniciar sesión y consultar sus calificaciones.
    </p>

    {% if session.get("usuario") %}
        <h3>Hola, {{ session["usuario"] }}</h3>
        <p>Ya puedes consultar tus notas en la sección Mis Notas.</p>
    {% else %}
        <p>Regístrate o inicia sesión para acceder al sistema.</p>
    {% endif %}

{% elif pagina == "registro" %}
    <h2>Registro de estudiantes</h2>

    <form method="POST">
        <input type="text" name="usuario" placeholder="Nombre de usuario" required>
        <input type="email" name="correo" placeholder="Correo electrónico" required>
        <input type="password" name="contrasena" placeholder="Contraseña (mínimo 8 caracteres)" required>
        <button type="submit">Crear cuenta</button>
    </form>

{% elif pagina == "login" %}
    <h2>Iniciar sesión</h2>

    <form method="POST">
        <input type="email" name="correo" placeholder="Correo electrónico" required>
        <input type="password" name="contrasena" placeholder="Contraseña" required>
        <button type="submit">Ingresar</button>
    </form>

    <p>
        <a href="/recuperar">¿Olvidaste tu contraseña?</a>
    </p>

{% elif pagina == "recuperar" %}
    <h2>Recuperar contraseña</h2>
    <p>Ingresa tus datos para actualizar tu contraseña.</p>

    <form method="POST">
        <input type="email" name="correo" placeholder="Correo registrado" required>
        <input type="text" name="codigo" placeholder="Código del administrador" required>
        <input type="password" name="nueva" placeholder="Nueva contraseña (mínimo 8 caracteres)" required>
        <button type="submit">Cambiar contraseña</button>
    </form>


{% elif pagina == "notas" %}
    <h2>Mis calificaciones</h2>
    <p>Estudiante: {{ session["usuario"] }}</p>

    {% if notas %}
        <table>
            <tr>
                <th>Materia</th>
                <th>Período 1</th>
                <th>Período 2</th>
                <th>Período 3</th>
            </tr>

            {% for nota in notas %}
            <tr>
                {% for dato in nota %}
                    <td>{{ dato if dato is not none else "-" }}</td>
                {% endfor %}
            </tr>
            {% endfor %}
        </table>
    {% else %}
        <p>Todavía no hay calificaciones registradas para tu cuenta.</p>
    {% endif %}
{% endif %}

</div>

<footer>
    <p>Instituto Técnico Nuevo Horizonte</p>
</footer>

</body>
</html>
"""


@app.route("/")
def inicio():
    return render_template_string(HTML, pagina="inicio")


@app.route("/registro", methods=["GET", "POST"])
def registro():
    preparar_excel()

    if request.method == "POST":
        usuario = request.form["usuario"].strip()
        correo = request.form["correo"].strip().lower()
        contrasena = request.form["contrasena"]

        if len(contrasena) < 8:
            flash("La contraseña debe tener al menos 8 caracteres.")
            return redirect("/registro")

        if buscar_usuario(correo):
            flash("Este correo ya está registrado.")
            return redirect("/registro")

        libro = load_workbook(ARCHIVO)
        hoja = libro["Estudiantes"]

        nuevo_id = hoja.max_row
        hoja.append([
            nuevo_id,
            usuario,
            correo,
            generate_password_hash(contrasena),
            datetime.now().strftime("%Y-%m-%d %H:%M")
        ])

        libro.save(ARCHIVO)
        libro.close()

        flash("Cuenta creada correctamente. Ahora puedes iniciar sesión.")
        return redirect("/login")

    return render_template_string(HTML, pagina="registro")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        correo = request.form["correo"].strip().lower()
        contrasena = request.form["contrasena"]

        usuario = buscar_usuario(correo)

        if usuario and check_password_hash(usuario[3], contrasena):
            session["usuario"] = usuario[1]
            session["correo"] = correo
            return redirect("/")

        flash("Correo o contraseña incorrectos.")
        return redirect("/login")

    return render_template_string(HTML, pagina="login")


@app.route("/notas")
def ver_notas():
    if not session.get("correo"):
        flash("Primero debes iniciar sesión.")
        return redirect("/login")

    notas = obtener_notas(session["correo"])

    return render_template_string(
        HTML,
        pagina="notas",
        notas=notas
    )

@app.route("/recuperar", methods=["GET", "POST"])
def recuperar():
    if request.method == "POST":
        correo = request.form["correo"].strip().lower()
        codigo = request.form["codigo"]
        nueva = request.form["nueva"]

        usuario = buscar_usuario(correo)

        if not usuario:
            flash("El correo no está registrado.")
        elif codigo != CODIGO_ADMIN:
            flash("El código del administrador es incorrecto.")
        elif len(nueva) < 8:
            flash("La contraseña debe tener al menos 8 caracteres.")
        else:
            libro = load_workbook(ARCHIVO)
            hoja = libro["Estudiantes"]

            for fila in range(2, hoja.max_row + 1):
                if hoja.cell(fila, 3).value.lower() == correo:
                    hoja.cell(fila, 4).value = generate_password_hash(nueva)
                    break

            libro.save(ARCHIVO)
            libro.close()

            flash("Contraseña actualizada. Ya puedes iniciar sesión.")
            return redirect("/login")

        return redirect("/recuperar")

    return render_template_string(HTML, pagina="recuperar")
@app.route("/salir")
def salir():
    session.clear()
    return redirect("/")


if __name__ == "__main__":
    preparar_excel()
    print("Sistema iniciado correctamente.")
    app.run(debug=True)
